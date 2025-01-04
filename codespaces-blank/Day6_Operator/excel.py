"""
from openpyxl import Workbook
from openpyxl.styles import PatternFill

workbook_load=Workbook()
sheet=workbook_load.active

data=[345,678,34,5,450]
for i, value in enumerate(data, start=1):
    cell=sheet.cell(row=i,column=1)
    cell.value=value
    if value<400:
        cell.fill=PatternFill(start_color="FF0000",end_color="FF0000",fill_type="solid")
workbook_load.save("excl.xlsx")
"""

"""
#open created workbook
from openpyxl import Workbook
from openpyxl.reader.excel import load_workbook

wb=load_workbook("excl.xlsx")
sheet=wb.active

print(sheet["A2"].value)
"""
#writing to excel files
"""
from openpyxl import Workbook

work_book = Workbook()
work_sheet = work_book.active

work_sheet.append(["student_name", "Roll_No", "Marks", "Remarks"])

students_data = [
    {"student_name": "John Doe", "Roll_No": 101, "Marks": 45},
    {"student_name": "Jane Smith", "Roll_No": 102, "Marks": 29},
    {"student_name": "Sam Wilson", "Roll_No": 103, "Marks": 78},
    {"student_name": "Anna Brown", "Roll_No": 104, "Marks": 32},
]

for i,student in enumerate(students_data, start=2):
    name = student["student_name"]
    roll_no = student["Roll_No"]
    marks = student["Marks"]
    remarks = f'=IF(C{i}>=33,"Pass","Fail")'
    work_sheet.append([name, roll_no, marks, remarks])

work_book.save("student_data.xlsx")

print("Excel file 'student_data.xlsx' has been created successfully.")
"""
""" #print the last column remarks for the above code
from openpyxl import Workbook

wb_load=load_workbook("student_data.xlsx")
sheet=wb_load.active

print("Remarks for each student:")
for cell in sheet["D"]:
    if cell.row >1:
        print(cell.value)
        """


from openpyxl import workbook
from openpyxl.reader.excel import load_workbook

wb=load_workbook("Std.xlsx")

for sheet_name in wb.sheetnames:
    print(sheet_name)
    sheet=wb[sheet_name]
    for row in sheet.iter_rows(min_row=1,max_col=1,max_row=4):
        for cell in row:
            print(cell.value)